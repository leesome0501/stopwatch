"""업로드된 파일을 시트별 셀 행렬(list[list[str]])로 읽는다.

은행 '엑셀 다운로드'는 실제로는 HTML 표인 .xls 가 흔하고(ADMIN 프로젝트에서 확인),
진짜 .xls(BIFF), .xlsx, CSV(UTF-8 또는 CP949)도 있다. 확장자 대신 파일 내용으로 판별한다.
"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime
from html.parser import HTMLParser

Matrix = list[list[str]]
Sheet = tuple[str, Matrix]   # (시트 이름, 셀 행렬)

XLSX_MAGIC = b"PK\x03\x04"
XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class UnreadableFile(ValueError):
    pass


def _cell_text(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S") if (v.hour or v.minute or v.second) else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v).strip()


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp949"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _looks_like_html(data: bytes) -> bool:
    head = data[:2048].lstrip().lower()
    return head.startswith((b"<html", b"<!doctype", b"<table", b"<meta", b"<?xml")) or b"<table" in head


class _TableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[Matrix] = []
        self._depth = 0
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._colspan = 1

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._depth += 1
            self.tables.append([])
        elif tag == "tr" and self._depth:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
            span = dict(attrs).get("colspan") or "1"
            self._colspan = int(span) if span.isdigit() else 1
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._row.extend([""] * (self._colspan - 1))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self.tables:
                self.tables[-1].append(self._row)
            self._row = None
        elif tag == "table" and self._depth:
            self._depth -= 1

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _read_html(data: bytes) -> list[Sheet]:
    p = _TableParser()
    p.feed(_decode(data))
    return [(f"table{i + 1}", t) for i, t in enumerate(p.tables) if t]


def _read_xlsx(data: bytes) -> list[Sheet]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [(ws.title, [[_cell_text(c) for c in row] for row in ws.iter_rows(values_only=True)])
                for ws in wb.worksheets]
    finally:
        wb.close()


def _read_xls(data: bytes) -> list[Sheet]:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    out: list[Sheet] = []
    for sh in book.sheets():
        rows: Matrix = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(_cell_text(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)))
                else:
                    row.append(_cell_text(cell.value))
            rows.append(row)
        out.append((sh.name, rows))
    return out


def _read_csv(data: bytes) -> list[Sheet]:
    text = _decode(data)
    first = text.splitlines()[0] if text else ""
    delim = "\t" if first.count("\t") > first.count(",") else ","
    return [("csv", [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delim)])]


def read_sheets(data: bytes) -> list[Sheet]:
    if not data:
        raise UnreadableFile("빈 파일입니다.")
    try:
        if data.startswith(XLSX_MAGIC):
            return _read_xlsx(data)
        if data.startswith(XLS_MAGIC):
            return _read_xls(data)
        if _looks_like_html(data):
            return _read_html(data)
        return _read_csv(data)
    except UnreadableFile:
        raise
    except Exception as e:  # 손상된 파일, 암호 걸린 엑셀 등
        raise UnreadableFile(f"파일을 읽을 수 없습니다: {e}") from e
