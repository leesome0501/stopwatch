"""AssetHub 표준 양식.

은행 형식으로 읽을 수 없는 자료(증권 매매, 배당, 환전, 부동산, 과거 내역 일괄 입력)를
사람이 채워 올리는 양식이다. GET /api/import-template 으로 내려받는다.
"""
from __future__ import annotations

import io
from decimal import Decimal

from .parsing import norm_header, parse_amount, parse_moment, parse_optional_amount
from .reader import Matrix
from .rows import ImportRow, ParseResult, RowError

COLUMNS = ["일자", "시간", "유형", "금액", "통화", "종목코드", "시장", "수량", "단가",
           "수수료", "세금", "카테고리", "거래처", "적요", "메모", "묶음", "거래후잔액"]
REQUIRED = {"일자", "유형"}

KIND_LABELS = {
    "입금": "cash_in", "출금": "cash_out", "매수": "buy", "매도": "sell",
    "배당": "dividend", "이자": "interest", "수수료": "fee", "세금": "tax",
    "환전": "fx", "입고": "instrument_in", "출고": "instrument_out",
    "분할": "split", "조정": "adjustment", "기초잔액": "adjustment",
}
KIND_LABELS.update({v: v for v in set(KIND_LABELS.values())})

# 금액 부호를 유형이 정하는 경우 (사용자는 양수로 써도 된다)
NEGATIVE_AMOUNT = {"cash_out", "buy", "fee", "tax"}
POSITIVE_AMOUNT = {"cash_in", "sell", "dividend", "interest"}
POSITIVE_QTY = {"buy", "instrument_in"}
NEGATIVE_QTY = {"sell", "instrument_out"}
NEEDS_INSTRUMENT = {"buy", "sell", "instrument_in", "instrument_out", "split"}

# 양식에 들어가지 않아야 할 시트 (예시·설명)
SKIP_SHEET_WORDS = ("예시", "설명", "안내")


def _header_index(matrix: Matrix) -> tuple[int, dict[str, int]] | None:
    for r, row in enumerate(matrix[:10]):
        h = [norm_header(c).rstrip("*") for c in row]
        if REQUIRED <= set(h) and "금액" in h:
            return r, {name: h.index(name) for name in COLUMNS if name in h}
    return None


def is_standard(matrix: Matrix) -> bool:
    return _header_index(matrix) is not None


def parse_standard_matrix(matrix: Matrix) -> ParseResult | None:
    found = _header_index(matrix)
    if found is None:
        return None
    hr, col = found
    result = ParseResult(format="AssetHub", rows=[])

    def get(row: list[str], name: str) -> str:
        i = col.get(name)
        return str(row[i]).strip() if i is not None and i < len(row) and row[i] is not None else ""

    for idx, row in enumerate(matrix[hr + 1:]):
        if not any(str(c).strip() for c in row):
            continue
        raw = {name: get(row, name) for name in col if get(row, name)}
        try:
            result.rows.append(_to_row(idx, get, row, raw))
        except ValueError as e:
            result.errors.append(RowError(idx, str(e), raw))
    return result


def _to_row(idx, get, row, raw) -> ImportRow:
    moment = parse_moment(get(row, "일자"), get(row, "시간"))
    if moment is None:
        raise ValueError(f"일자를 읽을 수 없습니다: {get(row, '일자')!r}")
    label = get(row, "유형")
    kind = KIND_LABELS.get(label) or KIND_LABELS.get(label.lower())
    if kind is None:
        raise ValueError(f"알 수 없는 유형입니다: {label!r} (가능: {', '.join(k for k in KIND_LABELS if not k.isascii())})")

    symbol = get(row, "종목코드") or None
    qty = parse_optional_amount(get(row, "수량"))
    price = parse_optional_amount(get(row, "단가"))
    fee = abs(parse_amount(get(row, "수수료")))
    tax = abs(parse_amount(get(row, "세금")))
    amount = parse_optional_amount(get(row, "금액"))
    group = get(row, "묶음") or None

    if kind in NEEDS_INSTRUMENT and (symbol is None or qty is None):
        raise ValueError(f"'{label}' 유형은 종목코드와 수량이 필요합니다.")
    if kind == "fx" and group is None:
        raise ValueError("환전은 두 줄(원화 출금, 외화 입금)을 같은 '묶음' 값으로 묶어야 합니다.")

    if qty is not None:
        if kind in POSITIVE_QTY:
            qty = abs(qty)
        elif kind in NEGATIVE_QTY:
            qty = -abs(qty)

    if amount is None:
        if kind in ("buy", "sell") and price is not None:
            gross = abs(qty) * price
            amount = -(gross + fee + tax) if kind == "buy" else gross - fee - tax
        elif kind in ("instrument_in", "instrument_out", "split"):
            amount = Decimal(0)
        else:
            raise ValueError(f"'{label}' 유형은 금액이 필요합니다.")
    if kind in NEGATIVE_AMOUNT:
        amount = -abs(amount)
    elif kind in POSITIVE_AMOUNT:
        amount = abs(amount)

    currency = (get(row, "통화") or "").upper() or None
    return ImportRow(
        file_index=idx, moment=moment, kind=kind, amount=amount,
        description=get(row, "적요"), counterparty=get(row, "거래처") or None,
        balance_after=parse_optional_amount(get(row, "거래후잔액")),
        currency=currency, instrument_symbol=symbol, instrument_market=get(row, "시장") or None,
        quantity=qty, unit_price=price, fee=fee, tax=tax,
        category_name=get(row, "카테고리") or None, memo=get(row, "메모") or None,
        group_label=group, raw=raw,
    )


EXAMPLES = [
    ["2026-10-01", "", "기초잔액", 3000000, "KRW", "", "", "", "", "", "", "", "", "기존 예수금", "", "", ""],
    ["2026-10-12", "09:05", "매수", "", "KRW", "005930", "KRX", 20, 60000, 150, "", "", "", "", "", "", ""],
    ["2026-10-13", "10:00", "환전", -1390000, "KRW", "", "", "", "", "", "", "", "", "달러 환전", "", "환전1", ""],
    ["2026-10-13", "10:00", "환전", 1000, "USD", "", "", "", "", "", "", "", "", "달러 환전", "", "환전1", ""],
    ["2026-10-13", "23:40", "매수", "", "USD", "AAPL", "NASDAQ", 4, 200, 1, "", "", "", "", "", "", ""],
    ["2026-10-15", "", "배당", 3.40, "USD", "AAPL", "NASDAQ", "", "", "", 0.60, "배당금", "", "세후 입금액", "", "", ""],
    ["2026-10-20", "", "매도", "", "KRW", "005930", "KRX", 5, 70000, 100, 1050, "", "", "", "", "", ""],
]

GUIDE = [
    ["열", "설명"],
    ["일자", "필수. 2026-10-01, 2026.10.01, 20261001 모두 가능"],
    ["시간", "선택. 같은 날 순서를 정할 때 사용"],
    ["유형", "필수. 입금, 출금, 매수, 매도, 배당, 이자, 수수료, 세금, 환전, 입고, 출고, 분할, 기초잔액"],
    ["금액", "계좌 현금 변화. 출금·매수·수수료·세금은 양수로 써도 자동으로 빠짐. 매수·매도는 비우면 수량×단가±수수료·세금으로 계산"],
    ["통화", "비우면 계좌 기본 통화"],
    ["종목코드/시장", "매수·매도·입고·출고·분할에 필수. 종목은 먼저 등록되어 있어야 함"],
    ["수량", "양수로 입력. 매도·출고는 자동으로 빠짐"],
    ["세금", "배당 원천징수, 매도 거래세. 금액에는 이미 반영된 값"],
    ["카테고리", "기본 카테고리 이름 (예: 급여, 배당금, 외식)"],
    ["묶음", "환전처럼 두 줄이 한 사건이면 같은 값을 적음"],
    ["거래후잔액", "선택. 적으면 잔액 검증과 중복 판정에 사용"],
]


def build_template() -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "거래"
    ws.append(COLUMNS)
    ex = wb.create_sheet("작성 예시")
    ex.append(COLUMNS)
    for r in EXAMPLES:
        ex.append(r)
    guide = wb.create_sheet("설명")
    for r in GUIDE:
        guide.append(r)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDEBF7")
    for sheet in (ws, ex, guide):
        for c in sheet[1]:
            c.font, c.fill = bold, fill
        for i, name in enumerate(COLUMNS if sheet is not guide else ["열", "설명"]):
            sheet.column_dimensions[chr(65 + i)].width = 14 if sheet is not guide else (14 if i == 0 else 90)
        sheet.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
